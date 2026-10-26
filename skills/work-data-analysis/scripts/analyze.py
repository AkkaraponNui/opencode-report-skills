"""Read local evidence and calculate explicitly defined service metrics.
Also supports profile-only inspection of unrelated tables; no guessed mappings.
"""
import argparse, csv, hashlib, json, math
from pathlib import Path
import pandas as pd
from docx import Document
from pypdf import PdfReader

FIELDS=['month','service','received','completed','on_time_completed','total_work_days','survey_responses','satisfaction_mean']
def write_json(p,value):
    p.write_text(json.dumps(value,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
def extract_sources(folder):
    result=[]
    for p in sorted(folder.iterdir()):
        if p.name.startswith('~$') or not p.is_file(): continue
        item={'file':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
        if p.suffix.lower()=='.docx':
            d=Document(p)
            item['segments']=[{'locator':f'paragraph {i}','text':t.text} for i,t in enumerate(d.paragraphs,1) if t.text.strip()]
            item['segments'] += [{'locator':f'table {i}','text':'\n'.join(' | '.join(c.text for c in r.cells) for r in t.rows)} for i,t in enumerate(d.tables,1)]
        elif p.suffix.lower()=='.pdf':
            item['segments']=[{'locator':f'page {i}','text':page.extract_text() or ''} for i,page in enumerate(PdfReader(p).pages,1)]
            if not any(x['text'].strip() for x in item['segments']): item['warning']='No extractable text; OCR required; do not invent contents.'
        elif p.suffix.lower() in ['.xlsx','.csv']:
            item['kind']='table'
        else: continue
        result.append(item)
    return result
def aggregate(frame):
    sums={c:float(frame[c].sum()) for c in FIELDS[2:7]}
    valid=frame['satisfaction_mean'].notna() & (frame['survey_responses']>0)
    denom=int(frame.loc[valid,'survey_responses'].sum())
    total=int(sums['survey_responses'])
    return {**{k:int(v) if v.is_integer() else v for k,v in sums.items()},
      'on_time_rate':sums['on_time_completed']/sums['completed'] if sums['completed'] else None,
      'average_days':sums['total_work_days']/sums['completed'] if sums['completed'] else None,
      'satisfaction':float((frame.loc[valid,'satisfaction_mean']*frame.loc[valid,'survey_responses']).sum()/denom) if denom else None,
      'scored_responses':denom,'score_coverage':denom/total if total else None}
def run(input_dir,output_dir,table='service-data.xlsx',sheet='Raw',mapping=None,profile_only=False,synthetic=False):
    src=Path(input_dir).resolve(); out=Path(output_dir).resolve()
    if out==src or src in out.parents: raise ValueError('Output must be outside input directory.')
    out.mkdir(parents=True,exist_ok=True)
    sources=extract_sources(src); write_json(out/'extracted_sources.json',sources)
    file=src/table
    frame=pd.read_csv(file) if file.suffix=='.csv' else pd.read_excel(file,sheet_name=sheet)
    if mapping: frame=frame.rename(columns=json.loads(Path(mapping).read_text(encoding='utf-8')))
    profile={'file':file.name,'sheet':sheet if file.suffix!='.csv' else None,'rows':len(frame),
             'columns':list(frame.columns),'missing':{str(c):int(frame[c].isna().sum()) for c in frame},
             'source_hash':hashlib.sha256(file.read_bytes()).hexdigest()}
    write_json(out/'profile.json',profile)
    if profile_only: return profile
    absent=set(FIELDS)-set(frame.columns)
    if absent: raise ValueError('Missing fields; use --profile-only or an explicit --mapping: '+','.join(sorted(absent)))
    frame=frame[FIELDS].copy()
    frame['_source_row']=range(2,len(frame)+2)
    if frame[['month','service']].isna().any().any(): raise ValueError('Missing month/service key.')
    if not frame['month'].astype(str).str.fullmatch(r'\d{4}-(0[1-9]|1[0-2])').all(): raise ValueError('month must use YYYY-MM.')
    for c in FIELDS[2:]:
        frame[c]=pd.to_numeric(frame[c],errors='raise')
        if not frame[c].dropna().map(math.isfinite).all(): raise ValueError('Non-finite value in '+c)
        if c!='satisfaction_mean' and frame[c].isna().any(): raise ValueError('Missing required measure '+c)
        if (frame[c].dropna()<0).any(): raise ValueError('Negative measure '+c)
    if ((frame['satisfaction_mean'].dropna()<1)|(frame['satisfaction_mean'].dropna()>5)).any(): raise ValueError('Satisfaction must be 1–5.')
    if (frame['on_time_completed']>frame['completed']).any(): raise ValueError('on_time_completed exceeds completed.')
    for c in ['received','completed','on_time_completed','survey_responses']:
        if (frame[c]%1!=0).any(): raise ValueError('Count must be integer: '+c)
    duplicates=frame.duplicated(subset=FIELDS,keep='first')
    duplicate_rows=frame.loc[duplicates,'_source_row'].astype(int).tolist()
    clean=frame.loc[~duplicates].copy()
    if clean.duplicated(['month','service']).any(): raise ValueError('Conflicting duplicate keys; resolve with data owner before aggregation.')
    missing_rows=clean.loc[clean['satisfaction_mean'].isna(),'_source_row'].astype(int).tolist()
    clean.to_csv(out/'cleaned.csv',index=False,encoding='utf-8-sig')
    monthly=[{'month':m,**aggregate(g),'source_rows':g['_source_row'].astype(int).tolist()} for m,g in clean.groupby('month',sort=True)]
    by_service=[{'service':s,**aggregate(g)} for s,g in clean.groupby('service',sort=True)]
    summary={'schema_version':1,'synthetic':synthetic,'source':file.name,'sheet':sheet,'source_sha256':profile['source_hash'],
      'period':[str(clean['month'].min()),str(clean['month'].max())],
      'quality':{'input_rows':len(frame),'used_rows':len(clean),'duplicate_rows_removed':duplicate_rows,'missing_score_rows':missing_rows},
      'totals':aggregate(clean),'monthly':monthly,'by_service':by_service,
      'methods':{'on_time_rate':'sum(on_time_completed)/sum(completed)',
       'average_days':'sum(total_work_days)/sum(completed)',
       'satisfaction':'sum(score*responses)/sum(responses), only rows with known scores',
       'received_vs_completed':'Different cohorts; do not treat completed/received as an individual completion rate.'},
      'limitations':['Aggregated data cannot establish causation.','Missing scores are excluded from score weights, never replaced with zero.']}
    write_json(out/'analysis.json',summary)
    pd.DataFrame(monthly).drop(columns='source_rows').to_csv(out/'monthly-metrics.csv',index=False,encoding='utf-8-sig')
    evidence=[{'metric':k,'value':v,'source':f'{file.name} / {sheet} / rows '+','.join(map(str,clean['_source_row'].astype(int))),
               'method':summary['methods'].get(k,'sum of included rows')} for k,v in summary['totals'].items()]
    write_json(out/'evidence.json',evidence)
    text=['# บันทึกคุณภาพและวิธีวิเคราะห์',f'- แถวข้อมูลเดิม {len(frame)} แถว ใช้จริง {len(clean)} แถว',
      f'- ตัดแถวซ้ำทุกช่องที่แถว Excel {duplicate_rows}; คงแถวแรกและไม่แก้ต้นฉบับ',
      f'- คะแนนที่ขาดที่แถว {missing_rows}; ไม่แทนด้วย 0',
      '- อัตราตรงเวลาใช้ผลรวมตัวตั้ง / ผลรวมตัวหาร ไม่เฉลี่ยเปอร์เซ็นต์รายแถว',
      '- คะแนนถ่วงน้ำหนักเฉพาะแถวที่มีคะแนน พร้อมแสดง coverage',
      '- เอกสารบริบทอ่านได้จาก extracted_sources.json; ข้อค้นพบต้องอ้างไฟล์และตำแหน่ง',
      '- ผลวิเคราะห์นี้เป็น descriptive statistics ไม่ได้พิสูจน์สาเหตุ']
    (out/'quality-notes.md').write_text('\n'.join(text),encoding='utf-8')
    return summary
if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--input',required=True);p.add_argument('--output',required=True)
    p.add_argument('--table',default='service-data.xlsx');p.add_argument('--sheet',default='Raw');p.add_argument('--mapping');p.add_argument('--profile-only',action='store_true');p.add_argument('--synthetic',action='store_true')
    a=p.parse_args();result=run(a.input,a.output,a.table,a.sheet,a.mapping,a.profile_only,a.synthetic)
    print(json.dumps(result.get('quality',result),ensure_ascii=True))
